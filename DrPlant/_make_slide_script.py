# -*- coding: utf-8 -*-
"""Tạo file Word: nội dung slide + văn nói thuyết trình DrPlant."""
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

OUT = Path(r"C:\AI-Agriculture-Ecosystem\DrPlant\NoiDung_ThuyetTrinh_DrPlant.docx")


def set_run_font(run, size=12, bold=False, italic=False, color=None):
    run.font.name = "Times New Roman"
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, 16, True, color=RGBColor(0x1F, 0x4E, 0x79))
    elif level == 1:
        r = p.add_run(text)
        set_run_font(r, 14, True, color=RGBColor(0x1F, 0x4E, 0x79))
    else:
        r = p.add_run(text)
        set_run_font(r, 12, True)
    p.paragraph_format.space_before = Pt(12 if level <= 1 else 6)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, bold=False, italic=False, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    r = p.add_run(text)
    set_run_font(r, 12, bold, italic)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(item)
        set_run_font(r, 12)


def add_speak(doc, text):
    add_body(doc, "Văn nói:", bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.3
    r = p.add_run(text)
    set_run_font(r, 12, False, True, color=RGBColor(0x33, 0x33, 0x33))


def add_note(doc, text):
    p = doc.add_paragraph()
    r = p.add_run("Ghi chú trình bày: ")
    set_run_font(r, 11, True, False, RGBColor(0x8B, 0x45, 0x00))
    r2 = p.add_run(text)
    set_run_font(r2, 11, False, True, RGBColor(0x8B, 0x45, 0x00))
    p.paragraph_format.space_after = Pt(6)


def hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("─" * 48)
    set_run_font(r, 9, False, False, RGBColor(0x99, 0x99, 0x99))


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2)

    # Cover of this script doc
    add_heading(doc, "NỘI DUNG SLIDE + VĂN NÓI THUYẾT TRÌNH", 0)
    add_heading(doc, "DrPlant – Ứng dụng Android hỗ trợ chẩn đoán bệnh lá cây", 0)
    add_body(doc, "Môn: Lập trình thiết bị di động  |  Trường: UEF  |  Năm: 2026", italic=True)
    add_body(
        doc,
        "Cách dùng: phần “Nội dung slide” để đưa lên PowerPoint/Canva; "
        "phần “Văn nói” để đọc khi thuyết trình (giọng tự nhiên, có thể rút ngắn nếu hết giờ).",
    )
    hr(doc)

    # ========== SLIDE 1 ==========
    add_heading(doc, "SLIDE 1 – Trang bìa / Mở đầu", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "Bộ Giáo dục và Đào tạo",
        "Trường Đại học Kinh tế – Tài chính TP. Hồ Chí Minh (UEF)",
        "Đề tài: DrPlant – Ứng dụng Android hỗ trợ chẩn đoán bệnh lá cây",
        "Môn học: Lập trình thiết bị di động",
        "Giảng viên hướng dẫn: Trần Thị Hường",
        "Thành viên nhóm:",
        "• Nguyễn Việt Anh – 235050236",
        "• Nguyễn Thế Anh – 235053028",
        "• Nguyễn Bảo Linh Nhi – 235057326",
        "• Phan Thị Ngọc Duyên – 235052977",
        "• Nguyễn Ngọc Phương Nghi – 235051516",
        "TP. Hồ Chí Minh – 2026",
    ])
    add_note(doc, "Logo trường / logo app DrPlant ở giữa hoặc góc trên. Nền sạch, chữ to, đủ đọc từ xa.")
    add_speak(
        doc,
        "Xin chào thầy cô và các bạn. Nhóm em xin trình bày đề tài DrPlant – "
        "một ứng dụng Android hỗ trợ người trồng nhận diện và theo dõi bệnh trên lá cây, "
        "trọng tâm là sầu riêng và cà phê. Giảng viên hướng dẫn của nhóm là cô Trần Thị Hường. "
        "Nhóm gồm năm thành viên: Việt Anh, Thế Anh, Linh Nhi, Ngọc Duyên và Phương Nghi. "
        "Trong phần sau, từng bạn sẽ giới thiệu ngắn phần việc mình phụ trách.",
    )
    hr(doc)

    # ========== SLIDE 2 ==========
    add_heading(doc, "SLIDE 2 – Mục lục", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "1. Giới thiệu thành viên & phân công",
        "2. Mô tả dự án & mục tiêu",
        "3. Mô hình dữ liệu (ERD / lưu trữ)",
        "4. Thiết kế giao diện (Figma / UI)",
        "5. Một số đoạn code chính",
        "6. Kiểm thử & test case",
        "7. Kết luận & hướng phát triển",
        "8. Q&A – Thank you",
    ])
    add_speak(
        doc,
        "Phần trình bày hôm nay đi theo tám mục chính. "
        "Trước hết nhóm giới thiệu thành viên và việc đã làm. "
        "Tiếp theo là mô tả dự án, mô hình dữ liệu, giao diện, "
        "một vài đoạn code then chốt, phần kiểm thử, rồi kết luận. "
        "Cuối cùng nhóm sẵn sàng nhận câu hỏi từ thầy cô.",
    )
    hr(doc)

    # ========== SLIDES 3–7 members ==========
    members = [
        (
            "SLIDE 3 – Thành viên: Nguyễn Việt Anh",
            [
                "Họ tên: Nguyễn Việt Anh",
                "MSSV: 235050236",
                "Vai trò: Lập trình chính (Android / Camera / Chẩn đoán)",
                "Công việc chính trong milestone:",
                "• CameraActivity: preview CameraX, chụp ảnh, flash, đảo camera",
                "• Luồng chọn ảnh từ Gallery",
                "• DiagnosisActivity – màn hình kết quả chẩn đoán",
                "• DiagnosisHistoryManager – lưu lịch sử trên máy",
                "• Cấu hình Gradle, tích hợp Firebase / ONNX Runtime",
            ],
            "Em là Nguyễn Việt Anh, phụ trách phần lập trình lõi trên Android. "
            "Em xử lý màn hình camera với CameraX, cho phép chụp hoặc chọn ảnh từ thư viện, "
            "rồi mở màn kết quả chẩn đoán. Phần lịch sử cũng do em gắn với SharedPreferences "
            "để lưu cục bộ trên máy. Ngoài ra em còn lo cấu hình dự án Gradle và kết nối "
            "Firebase cùng engine ONNX chạy trên thiết bị.",
        ),
        (
            "SLIDE 4 – Thành viên: Nguyễn Thế Anh",
            [
                "Họ tên: Nguyễn Thế Anh",
                "MSSV: 235053028",
                "Vai trò: Phân tích yêu cầu / UI thống kê / điều hướng",
                "Công việc chính:",
                "• Ghi nhận yêu cầu chức năng & use case",
                "• StatisticalActivity – giao diện tổng quan thống kê",
                "• Điều hướng giữa các màn hình chính",
                "• Chuẩn hóa chuỗi giao diện (strings.xml) và tài nguyên UI",
            ],
            "Em là Nguyễn Thế Anh. Em tập trung vào phần phân tích: viết yêu cầu, "
            "mô tả use case cho nhóm dễ chia việc. Trên app, em làm màn thống kê "
            "StatisticalActivity và giúp nối luồng điều hướng giữa các màn. "
            "Em cũng chỉnh các chuỗi tiếng Việt trong resources cho giao diện thống nhất.",
        ),
        (
            "SLIDE 5 – Thành viên: Nguyễn Bảo Linh Nhi",
            [
                "Họ tên: Nguyễn Bảo Linh Nhi",
                "MSSV: 235057326",
                "Vai trò: Kiểm thử UI trên thiết bị / lịch sử / minh họa báo cáo",
                "Công việc chính:",
                "• Rà soát UI và test thực tế trên máy / emulator",
                "• Kiểm tra bộ lọc lịch sử (loại cây, ngày)",
                "• Chụp screenshot các màn hình phục vụ báo cáo & slide",
                "• Góp ý trải nghiệm người dùng (UX)",
            ],
            "Em là Nguyễn Bảo Linh Nhi. Em đảm nhận kiểm thử trên thiết bị thật và emulator: "
            "xem layout có lỗi không, bộ lọc lịch sử chạy đúng chưa. "
            "Em cũng chụp lại các màn hình để đưa vào báo cáo tiến độ và phần minh họa trên slide, "
            "đồng thời góp ý để giao diện dễ dùng hơn với người trồng cây.",
        ),
        (
            "SLIDE 6 – Thành viên: Phan Thị Ngọc Duyên",
            [
                "Họ tên: Phan Thị Ngọc Duyên",
                "MSSV: 235052977",
                "Vai trò: Authentication UI & nội dung tài liệu",
                "Công việc chính:",
                "• LoginActivity / RegisterActivity và layout XML",
                "• Luồng đăng nhập, đăng ký, quên mật khẩu (Firebase Auth)",
                "• Rà soát bảng biểu trong báo cáo",
                "• Ghi nhận yêu cầu phi chức năng (hiệu năng, quyền, UX)",
            ],
            "Em là Phan Thị Ngọc Duyên. Em làm hai màn đăng nhập và đăng ký, "
            "kèm layout XML và nối với Firebase Authentication: đăng nhập email–mật khẩu, "
            "đăng ký tài khoản mới và gửi email đặt lại mật khẩu. "
            "Em còn hỗ trợ kiểm tra bảng trong báo cáo và liệt kê các yêu cầu phi chức năng "
            "như quyền camera, phản hồi UI khi mạng yếu.",
        ),
        (
            "SLIDE 7 – Thành viên: Nguyễn Ngọc Phương Nghi",
            [
                "Họ tên: Nguyễn Ngọc Phương Nghi",
                "MSSV: 235051516",
                "Vai trò: Quyền truy cập / lưu trữ cục bộ / theo dõi tiến độ",
                "Công việc chính:",
                "• Checklist quyền CAMERA & đọc ảnh theo phiên bản Android",
                "• Ghi chú thiết kế lưu trữ cục bộ (SharedPreferences + Gson)",
                "• Theo dõi tiến độ module so với kế hoạch",
                "• Hỗ trợ chuẩn bị tài liệu thuyết trình",
            ],
            "Em là Nguyễn Ngọc Phương Nghi. Em phụ trách phần quyền trên Android: "
            "camera và đọc ảnh phải xin đúng cách trên các bản Android mới. "
            "Em còn ghi chú cách app lưu dữ liệu cục bộ bằng SharedPreferences, "
            "và theo dõi tiến độ từng module so với kế hoạch để nhóm biết phần nào xong, "
            "phần nào còn dang dở như thống kê và đồng bộ server.",
        ),
    ]

    for title, bullets, speak in members:
        add_heading(doc, title, 1)
        add_body(doc, "Nội dung slide:", bold=True)
        add_bullets(doc, bullets)
        add_note(doc, "Có thể thêm ảnh chân dung nhỏ hoặc icon vai trò. Mỗi slide thành viên khoảng 30–45 giây.")
        add_speak(doc, speak)
        hr(doc)

    # ========== SLIDE 8 ==========
    add_heading(doc, "SLIDE 8 – Mô tả dự án & mục tiêu", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_body(doc, "Mô tả ngắn:", bold=True)
    add_bullets(doc, [
        "DrPlant giúp người trồng chụp / chọn ảnh lá cây → nhận kết quả gợi ý trên máy",
        "Đối tượng: lá sầu riêng, cà phê (mở rộng sau)",
        "Nền tảng: Android (Java), CameraX, Firebase Auth, ONNX Runtime on-device",
        "Không cần gửi ảnh lên server trong milestone hiện tại (chẩn đoán chạy trên máy)",
    ])
    add_body(doc, "Mục tiêu:", bold=True)
    add_bullets(doc, [
        "Đăng ký / đăng nhập an toàn bằng Firebase",
        "Chụp hoặc chọn ảnh lá → xem kết quả chẩn đoán",
        "Lưu & xem lại lịch sử trên thiết bị",
        "Xây khung màn thống kê (dữ liệu thật sẽ bổ sung giai đoạn sau)",
    ])
    add_note(
        doc,
        "Chèn 2–3 ảnh minh họa: logo app, màn Login, màn Camera (file sẵn trong "
        "report_figures/ui/). Slide gọn – tránh đoạn văn dài.",
    )
    add_speak(
        doc,
        "DrPlant ra đời từ thực tế: nhiều nhà vườn vẫn nhìn lá bằng mắt thường, "
        "khó theo dõi bệnh theo thời gian. App cho phép đăng nhập, chụp hoặc chọn ảnh lá, "
        "rồi xem kết quả gợi ý ngay trên điện thoại. "
        "Hiện nhóm tập trung hai loại cây là sầu riêng và cà phê. "
        "Phần xác thực dùng Firebase; phần nhận diện chạy mô hình ONNX ngay trên máy "
        "nên không phụ thuộc gửi ảnh lên server. "
        "Mục tiêu milestone này là hoàn thiện luồng chính: tài khoản, camera, kết quả, "
        "lịch sử; còn thống kê theo khu vực và đồng bộ cloud sẽ làm tiếp.",
    )
    hr(doc)

    # ========== SLIDE 9 ERD ==========
    add_heading(doc, "SLIDE 9 – Sơ đồ ERD / Mô hình dữ liệu", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "Giai đoạn hiện tại: chưa dùng Room / SQL server dùng chung",
        "Thực thể logic đề xuất (hướng triển khai đầy đủ):",
        "• User (uid, email) – quản lý bởi Firebase Auth",
        "• DiagnosisHistoryItem (id, viewedAt, imagePath, plantType, status, note…)",
        "• Area / Plant / Disease (dự kiến – phục vụ thống kê; chưa CRUD)",
        "Quan hệ: User 1—N HistoryItem; Area 1—N Plant; Plant N—N Disease (qua bản ghi bệnh)",
        "Lưu tạm trên máy: SharedPreferences + JSON (Gson), tối đa ~100 mục lịch sử",
    ])
    add_body(doc, "Gợi ý vẽ ERD trên slide:", bold=True)
    add_bullets(doc, [
        "User ||--o{ DiagnosisHistoryItem",
        "Area ||--o{ Plant",
        "Plant }o--o{ DiseaseRecord",
        "Ghi chú góc slide: “Phần Area/Plant/Disease: thiết kế trước, implement sau”",
    ])
    add_note(
        doc,
        "Vẽ ERD bằng draw.io / PowerPoint. Nếu giảng viên hỏi DB: giải thích rõ milestone "
        "ưu tiên on-device; ERD là định hướng kiến trúc dữ liệu.",
    )
    add_speak(
        doc,
        "Về dữ liệu, nhóm tách hai lớp. Thứ nhất, tài khoản người dùng do Firebase Auth giữ – "
        "app chỉ nhận uid và email sau khi đăng nhập. "
        "Thứ hai, mỗi lần chẩn đoán được lưu thành một mục lịch sử trên máy: "
        "có id, thời điểm xem, đường dẫn ảnh, loại cây, trạng thái và ghi chú. "
        "Nhóm cũng phác ERD cho khu vực, cây trồng và bệnh để phục vụ thống kê sau này, "
        "nhưng milestone này chưa tạo database Room hay server SQL. "
        "Hiện lịch sử được serialize JSON trong SharedPreferences cho nhanh và đủ dùng trên mobile.",
    )
    hr(doc)

    # ========== SLIDE 10 Figma ==========
    add_heading(doc, "SLIDE 10 – Thiết kế giao diện (Figma / UI)", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "Nguyên tắc: màn camera tối giản – ưu tiên khung hình; form đăng nhập rõ ràng",
        "Các màn chính:",
        "• Login / Register – form cuộn, màu nhấn tím/primary Material",
        "• Camera (Home) – preview full, nút chụp, gallery, menu lịch sử / thống kê / đăng xuất",
        "• Diagnosis – hiển thị kết quả + nút quay lại / xem lịch sử",
        "• History – lọc theo loại cây & ngày",
        "• Statistics – dashboard tổng quan (UI, dữ liệu thật giai đoạn sau)",
        "Link Figma (điền khi có): …………………………",
    ])
    add_note(
        doc,
        "Nếu chưa có file Figma: chèn screenshot thật từ emulator (ui_login, ui_camera, "
        "ui_history…) và nói “prototype đã hiện thực trực tiếp trên Android XML”. "
        "Có thể export frame Figma sau rồi dán link lên slide.",
    )
    add_speak(
        doc,
        "Về giao diện, nhóm thiết kế theo hướng dễ dùng ngoài đồng: "
        "màn đăng nhập gọn, còn màn camera để preview chiếm phần lớn màn hình. "
        "Sau khi có kết quả, người dùng xem Diagnosis rồi có thể mở History để lọc theo "
        "sầu riêng, cà phê hoặc theo ngày. Màn thống kê đã dựng layout tổng quan "
        "nhưng chưa gắn số liệu thật vì thiếu dữ liệu khu vực và cây trồng. "
        "Nhóm sẽ để link Figma trên slide; phần demo hôm nay dùng đúng giao diện đã build trong APK.",
    )
    hr(doc)

    # ========== SLIDE 11 Code Create DB / Local ==========
    add_heading(doc, "SLIDE 11 – Code chính: Lưu trữ cục bộ (thay Create DB)", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "File: DiagnosisHistoryManager.java",
        "Ý tưởng: SharedPreferences + Gson ↔ List<DiagnosisHistoryItem>",
        "Đoạn code minh họa (rút gọn trên slide):",
    ])
    code1 = (
        "SharedPreferences prefs = context.getSharedPreferences(\n"
        "    \"diagnosis_history_prefs\", Context.MODE_PRIVATE);\n"
        "String json = prefs.getString(\"history_items\", \"\");\n"
        "List<DiagnosisHistoryItem> items =\n"
        "    new Gson().fromJson(json, HISTORY_LIST_TYPE);\n"
        "// saveHistoryItem(...): thêm phần tử mới, giới hạn 100 mục"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(code1)
    set_run_font(r, 10, False, False, RGBColor(0x1A, 0x1A, 0x2E))
    add_note(doc, "Nói rõ: milestone này chưa Create SQL DB; đây là lớp lưu trữ tương đương cho mobile offline.")
    add_speak(
        doc,
        "Thay vì tạo database SQL ngay từ đầu, nhóm dùng SharedPreferences để lưu danh sách "
        "lịch sử dưới dạng JSON. Mỗi lần chẩn đoán xong, app thêm một phần tử vào đầu list, "
        "giới hạn khoảng một trăm mục để tránh phình dữ liệu. "
        "Cách này đủ cho demo offline và khớp với phạm vi môn học; "
        "khi làm Area và Plant đầy đủ, nhóm sẽ chuyển sang Room hoặc backend.",
    )
    hr(doc)

    # ========== SLIDE 12 Auth ==========
    add_heading(doc, "SLIDE 12 – Code chính: Authentication", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "File: LoginActivity.java / RegisterActivity.java",
        "Thư viện: Firebase Authentication (email / password)",
        "Luồng: validate → signInWithEmailAndPassword → CameraActivity",
        "Đoạn code minh họa:",
    ])
    code2 = (
        "if (!isValidEmail(email)) { emailLayout.setError(...); return; }\n"
        "if (TextUtils.isEmpty(password)) { ...; return; }\n"
        "firebaseAuth.signInWithEmailAndPassword(email, password)\n"
        "    .addOnCompleteListener(this, task -> {\n"
        "        if (task.isSuccessful()) goToCamera();\n"
        "        else showLoginError(task.getException());\n"
        "    });"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(code2)
    set_run_font(r, 10, False, False, RGBColor(0x1A, 0x1A, 0x2E))
    add_bullets(doc, [
        "Bổ sung: sendPasswordResetEmail khi quên mật khẩu",
        "onStart: nếu đã đăng nhập → vào thẳng CameraActivity",
    ])
    add_speak(
        doc,
        "Phần đăng nhập, nhóm kiểm tra email đúng định dạng và mật khẩu không để trống "
        "trước khi gọi Firebase. Nếu đăng nhập thành công thì chuyển sang CameraActivity; "
        "nếu sai thì hiện thông báo lỗi rõ ràng. "
        "Người dùng cũng có thể nhận email đặt lại mật khẩu. "
        "Khi mở lại app mà phiên vẫn còn, onStart sẽ đưa thẳng vào màn chính, khỏi đăng nhập lại.",
    )
    hr(doc)

    # ========== SLIDE 13 Security ==========
    add_heading(doc, "SLIDE 13 – Code chính: Security (Bảo mật)", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "Xác thực phía Firebase – không tự lưu mật khẩu dạng plain text trên máy",
        "Runtime permission: CAMERA, READ_MEDIA_IMAGES (Android 13+)",
        "Activity không cần export thì android:exported=\"false\"",
        "Validate input trước khi gọi API Auth (tránh request thừa / lỗi khó hiểu)",
        "google-services.json & SHA-1 gắn với project Firebase của nhóm (không public key)",
        "Mật khẩu đăng ký tối thiểu 6 ký tự (theo quy ước Firebase)",
    ])
    add_body(doc, "Snippet quyền (ý tưởng):", bold=True)
    code3 = (
        "requestPermission.launch(Manifest.permission.CAMERA);\n"
        "// Chỉ mở preview CameraX sau khi được cấp quyền"
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.5)
    r = p.add_run(code3)
    set_run_font(r, 10, False, False, RGBColor(0x1A, 0x1A, 0x2E))
    add_speak(
        doc,
        "Về bảo mật, nhóm không tự chứa mật khẩu trên máy mà giao cho Firebase Auth. "
        "Camera và ảnh chỉ dùng được sau khi người dùng cấp quyền lúc chạy. "
        "Các Activity nội bộ được để exported bằng false để app khác không gọi lung tung. "
        "File cấu hình Firebase và chữ ký SHA giữ trong nhóm, không đưa lên slide chi tiết. "
        "Những bước này đủ để demo an toàn ở mức ứng dụng học phần.",
    )
    hr(doc)

    # ========== SLIDE 14 Test ==========
    add_heading(doc, "SLIDE 14 – Test lỗi / Test case", 1)
    add_body(doc, "Nội dung slide (bảng rút gọn):", bold=True)

    table = doc.add_table(rows=7, cols=4)
    table.style = "Table Grid"
    headers = ["Mã", "Tình huống", "Kỳ vọng", "Kết quả"]
    data = [
        ["TC01", "Đăng nhập sai mật khẩu", "Báo lỗi, không vào Camera", "Pass"],
        ["TC02", "Email sai định dạng", "Hiện lỗi trên ô Email", "Pass"],
        ["TC03", "Từ chối quyền Camera", "Không mở preview / có hướng dẫn", "Pass"],
        ["TC04", "Chụp / chọn ảnh → Diagnosis", "Mở màn kết quả", "Pass"],
        ["TC05", "Lọc History theo cà phê", "Chỉ hiện mục cà phê", "Pass"],
        ["TC06", "Màn Statistics chưa có data", "UI hiện placeholder “—” / chưa có dữ liệu", "Pass*"],
    ]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(h)
        set_run_font(r, 11, True)
    for r_i, row in enumerate(data, start=1):
        for c_i, val in enumerate(row):
            table.rows[r_i].cells[c_i].text = ""
            rr = table.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
            set_run_font(rr, 10)

    add_body(doc, "")
    add_bullets(doc, [
        "Môi trường test: Android Emulator (Medium Phone) + thiết bị thật (nếu có)",
        "Pass*: đúng với phạm vi hiện tại (UI only, chưa gắn dữ liệu khu vực)",
        "Lỗi đã gặp & xử lý: dialog tương thích 16KB trên emulator mới; thiếu google-services lúc đầu",
    ])
    add_speak(
        doc,
        "Nhóm kiểm thử theo vài tình huống thực tế. "
        "Ví dụ đăng nhập sai mật khẩu phải báo lỗi và không cho vào camera; "
        "email sai định dạng thì hiện ngay trên ô nhập. "
        "Nếu người dùng từ chối quyền camera thì app không cố mở preview. "
        "Luồng vui: chụp hoặc chọn ảnh rồi sang màn kết quả; lịch sử lọc theo loại cây. "
        "Màn thống kê cố ý để trống số liệu – test case ghi nhận là đúng với tiến độ hiện tại. "
        "Trong quá trình dựng, nhóm cũng xử lý cấu hình Firebase và cảnh báo tương thích emulator.",
    )
    hr(doc)

    # ========== SLIDE 15 Conclusion ==========
    add_heading(doc, "SLIDE 15 – Kết luận", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_body(doc, "Đã hoàn thành:", bold=True)
    add_bullets(doc, [
        "Auth Firebase: đăng ký, đăng nhập, đăng xuất, quên mật khẩu",
        "CameraX + Gallery → màn kết quả chẩn đoán",
        "Lịch sử cục bộ + bộ lọc",
        "Khung UI thống kê",
        "Báo cáo tiến độ & bộ screenshot minh họa",
    ])
    add_body(doc, "Hạn chế / việc tiếp theo:", bold=True)
    add_bullets(doc, [
        "Statistics chưa có dữ liệu Area / Plant thật",
        "Chưa đồng bộ server / database dùng chung",
        "Cải thiện độ chính xác mô hình ONNX & trải nghiệm offline",
        "Bổ sung Figma hoàn chỉnh + test case mở rộng",
    ])
    add_speak(
        doc,
        "Tóm lại, milestone này nhóm đã dựng được phần vỏ Android chạy được: "
        "có tài khoản, camera, kết quả, lịch sử và khung thống kê. "
        "Phần còn lại rõ ràng: gắn dữ liệu khu vực–cây trồng để thống kê có ý nghĩa, "
        "rồi tính tới đồng bộ server. "
        "Nhóm xem đây là nền vững để phát triển tiếp trong các giai đoạn sau của đề tài.",
    )
    hr(doc)

    # ========== SLIDE 16 Thank you ==========
    add_heading(doc, "SLIDE 16 – Thank you", 1)
    add_body(doc, "Nội dung slide:", bold=True)
    add_bullets(doc, [
        "Thank you / Xin cảm ơn thầy cô và các bạn",
        "DrPlant – Android Application for Plant Leaf Diagnosis Support",
        "Q&A – Nhóm sẵn sàng nhận câu hỏi",
        "(Tuỳ chọn) Email liên hệ nhóm / GitHub repo",
    ])
    add_note(doc, "Nền tối giản, chữ Thank you lớn; có thể để QR GitHub nếu được phép.")
    add_speak(
        doc,
        "Xin cảm ơn thầy cô và mọi người đã lắng nghe. "
        "Nhóm DrPlant rất mong nhận góp ý để hoàn thiện ứng dụng. "
        "Em xin phép dừng phần trình bày và sẵn sàng trả lời câu hỏi.",
    )
    hr(doc)

    # Appendix tips
    add_heading(doc, "PHỤ LỤC – Gợi ý chia người nói & thời lượng", 1)
    add_bullets(doc, [
        "Tổng ~10–12 phút: Slide 1–2 (30s) → mỗi thành viên 30–40s → mô tả dự án 1 phút "
        "→ ERD 1 phút → UI 1 phút → 3 slide code ~2.5 phút → test 1 phút → kết luận + thank you 1 phút",
        "Nên demo APK live 30–60 giây giữa slide UI và code (đăng nhập → camera).",
        "Nếu thiếu giờ: gộp 5 slide thành viên thành 1 slide “Nhóm & phân công”, "
        "văn nói từng bạn vẫn giữ nhưng nói nhanh hơn.",
    ])

    add_heading(doc, "Checklist đáp ứng yêu cầu đề bài slide", 1)
    add_bullets(doc, [
        "✓ Slide đầu: tên đề tài, GVHD, thành viên nhóm",
        "✓ Slide mục lục",
        "✓ Slide giới thiệu từng thành viên + task",
        "✓ Mô tả dự án, mục tiêu (ngắn gọn + gợi ý ảnh minh họa)",
        "✓ Sơ đồ ERD / mô hình dữ liệu",
        "✓ Figma / UI",
        "✓ Code chính: lưu trữ (Create DB tương đương), Authentication, Security",
        "✓ Test lỗi / test case",
        "✓ Kết luận",
        "✓ Thank you",
        "✓ Văn nói tự nhiên kèm từng slide",
    ])

    doc.save(str(OUT))
    # mirror copy
    alt = OUT.with_name("NoiDung_ThuyetTrinh_DrPlant_updated.docx")
    doc.save(str(alt))
    print("Saved", OUT)
    print("Saved", alt)


if __name__ == "__main__":
    main()
